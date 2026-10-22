#!/usr/bin/env python3
"""Generate PPT from docs/机房动环管理系统-四台服务器使用方案-简版.md structure."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "机房动环管理系统-四台服务器使用方案-简版.pptx"

NAVY = RGBColor(0x06, 0x5A, 0x82)
TEAL = RGBColor(0x02, 0x80, 0x90)
SLATE = RGBColor(0x33, 0x41, 0x55)
MUTED = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF1, 0xF5, 0xF9)


def set_run(run, size=18, bold=False, color=SLATE, name="PingFang SC"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name


def add_title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, Inches(2.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = NAVY
    shape.line.fill.background()

    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.55), Inches(11.5), Inches(1.2))
    tf = box.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "机房动环管理系统"
    set_run(r, 36, True, WHITE)
    p = tf.add_paragraph()
    r = p.add_run()
    r.text = "落地实施安排（四台服务器 · 808 基地）"
    set_run(r, 22, False, RGBColor(0xCA, 0xDC, 0xFC))

    sub = slide.shapes.add_textbox(Inches(0.6), Inches(2.6), Inches(11), Inches(1))
    tf = sub.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "完成节点：2026 年 12 月 10 日"
    set_run(r, 20, True, TEAL)
    p = tf.add_paragraph()
    r = p.add_run()
    r.text = "编制单位：中移集成"
    set_run(r, 16, False, MUTED)


def add_bullets_slide(prs, title, bullets, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(0.12), prs.slide_height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = TEAL
    bar.line.fill.background()

    tb = slide.shapes.add_textbox(Inches(0.45), Inches(0.35), Inches(12), Inches(0.8))
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = title
    set_run(r, 28, True, NAVY)

    if subtitle:
        st = slide.shapes.add_textbox(Inches(0.45), Inches(1.05), Inches(12), Inches(0.5))
        r = st.text_frame.paragraphs[0].add_run()
        r.text = subtitle
        set_run(r, 14, False, MUTED)

    body_top = 1.45 if subtitle else 1.15
    body = slide.shapes.add_textbox(Inches(0.55), Inches(body_top), Inches(12.2), Inches(5.8))
    tf = body.text_frame
    tf.word_wrap = True
    for i, line in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = 0
        p.space_after = Pt(10)
        r = p.add_run()
        r.text = line
        set_run(r, 17, False, SLATE)


def add_table_slide(prs, title, headers, rows, font_size=11):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Inches(0.45), Inches(0.25), Inches(12), Inches(0.6))
    r = tb.text_frame.paragraphs[0].add_run()
    r.text = title
    set_run(r, 24, True, NAVY)

    nrows = len(rows) + 1
    ncols = len(headers)
    left, top, width, height = Inches(0.35), Inches(0.95), Inches(12.6), Inches(6.0)
    table = slide.shapes.add_table(nrows, ncols, left, top, width, height).table

    col_widths = [1.0, 1.35, 2.0, 1.0] if ncols == 4 else None
    if col_widths and ncols == 4:
        for i, w in enumerate(col_widths):
            table.columns[i].width = Inches(w)

    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = h
        set_run(r, font_size, True, WHITE)

    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            if i % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j > 0 else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = val
            set_run(r, font_size, j == 0 or (ncols == 4 and j == 1), SLATE)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(prs)

    add_bullets_slide(
        prs,
        "总体安排",
        [
            "本方案计划将动环管理系统软件需求讨论、定制开发及部署测试等工作以 808 大楼为基地。",
            "前期大部分工作远程开发（自有环境，不接客户生产数据）。",
            "机房架构、机柜数量等资料脱敏后由甲方按需提供。",
            "需到现场时：以 808 为基地，与甲方进行需求沟通、方案讨论、系统调试、演示、试用等。",
        ],
    )

    add_table_slide(
        prs,
        "一、四台服务器分工（现位于 808）",
        ["机器", "用途", "说明"],
        [
            ["1 号", "设备接入与采集调试", "非日常登录入口"],
            ["2 号", "主系统（浏览器、大屏、日常监测）", "客户唯一常用入口"],
            ["3 号", "同 2 号版本，升级与验收演练", "平时演示仍用 2 号"],
            ["4 号", "备份与安装升级资料", "不对业务开放"],
        ],
        font_size=14,
    )

    add_bullets_slide(
        prs,
        "服务器部署说明",
        [
            "四台机目前在 808 安装部署和运行。",
            "仅当实际业务需要的时候再进行调整。",
        ],
    )

    phases = [
        ["需求与方案", "10/20～10/26", "需求对接、方案与计划确认", "远程为主；808 基地短访"],
        ["功能开发", "10/27～11/09", "按甲方要求定制功能（模拟数据）", "远程为主；808 按需"],
        ["808 试装演示", "11/10～11/16", "版本部署、演示与问题收集", "808 基地；启用四台"],
        ["对接与完善", "11/17～11/30", "设备数据接入；308 与 808 关联", "808 基地；308 按需"],
        ["联合测试", "12/01～12/07", "全链路测试、整改、版本冻结", "808 基地"],
        ["正式交付", "12/10", "按确认范围交付使用", "808 正式使用（2 号主系统）"],
    ]
    add_table_slide(
        prs,
        "二、分阶段落地计划",
        ["阶段", "时间", "主要工作", "工作地点"],
        phases,
        font_size=12,
    )

    add_bullets_slide(
        prs,
        "客户配合与说明",
        [
            "808 驻场前提前 1～2 个工作日提供人员信息备案。",
            "阶段切换以计划表为准；有调整由中移集成书面同步。",
            "编制单位：中移集成",
        ],
        subtitle="前期开发/部署环境：自有环境或 808 四台；客户四台在试装演示阶段正式启用。",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
